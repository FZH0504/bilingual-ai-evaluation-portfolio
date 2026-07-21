"""Create input and optional external-metric Excel templates."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent / "data"

COLUMNS = ["text_id", "condition", "model_name", "model_version_or_access_label",
           "access_date", "prompt_id", "prompt_text", "source_zh", "reference_en",
           "output_en", "generation_id", "notes"]
DESCRIPTIONS = {
    "text_id": "完整文本的配对标识；同一文本在两个条件下必须一致。",
    "condition": "提示词条件名称；可自由配置，不写死。",
    "model_name": "生成译文的模型名称。",
    "model_version_or_access_label": "模型版本、界面标签或访问标签。",
    "access_date": "模型访问/生成日期。",
    "prompt_id": "提示词标识；应与 condition 的实验设计一致。",
    "prompt_text": "原始提示词全文。",
    "source_zh": "中文原文（COMET src），不得润色。",
    "reference_en": "人工英文参考译文（COMET ref），不得润色。",
    "output_en": "模型英文译文（COMET mt），不得润色。",
    "generation_id": "重复生成标识，默认 run_1。",
    "notes": "可选备注。",
}
SEGMENT_COLUMNS = ["text_id", "condition", "segment_id", "source_segment",
                   "reference_segment", "output_segment"]
EXTERNAL_COLUMNS = ["text_id", "condition", "tool_name", "tool_version_or_access_date",
                    "syntactic_simplicity", "word_concreteness", "referential_cohesion",
                    "deep_cohesion", "l2_readability",
                    "syntactic_simplicity_reference", "word_concreteness_reference",
                    "referential_cohesion_reference", "deep_cohesion_reference",
                    "l2_readability_reference"]


def format_workbook(path: Path) -> None:
    """Apply readable widths, filters, frozen rows, and header styling."""
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for index, column in enumerate(worksheet.columns, start=1):
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
            worksheet.column_dimensions[get_column_letter(index)].width = max(width, 12)
    workbook.save(path)


def main() -> None:
    """Write input_template.xlsx and external_metrics.xlsx without overwriting data."""
    DATA.mkdir(parents=True, exist_ok=True)
    template = DATA / "input_template.xlsx"
    with pd.ExcelWriter(template, engine="openpyxl") as writer:
        pd.DataFrame(columns=COLUMNS).to_excel(writer, sheet_name="data", index=False)
        pd.DataFrame({"column": COLUMNS, "purpose": [DESCRIPTIONS[c] for c in COLUMNS]}).to_excel(
            writer, sheet_name="instructions", index=False)
        pd.DataFrame(columns=SEGMENT_COLUMNS).to_excel(writer, sheet_name="aligned_segments", index=False)
    format_workbook(template)
    external = DATA / "external_metrics.xlsx"
    if not external.exists():
        pd.DataFrame(columns=EXTERNAL_COLUMNS).to_excel(external, index=False, sheet_name="external_metrics")
        format_workbook(external)
    print(f"Created: {template}")
    print("Copy it to data/input.xlsx after filling it; the original template is preserved.")


if __name__ == "__main__":
    main()
